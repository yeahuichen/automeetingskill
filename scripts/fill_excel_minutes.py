# -*- coding: utf-8 -*-
"""
标准会议纪要 Excel 安全无损回填工具
严格按照《软件研发部产品周例会会议纪要》模板标准，向表格无损写入决议与待办项。
"""

import os
import sys
import json
import argparse
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

def fill_minutes_excel(excel_path, items, resolutions_text=None, output_path=None):
    excel_path = os.path.abspath(excel_path)
    if not os.path.isfile(excel_path):
        raise FileNotFoundError(f"未找到 Excel 模板文件: {excel_path}")
        
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    
    # 1. 如果提供了决议文本，安全更新决议部分 (A11 单元格)
    if resolutions_text:
        ws.cell(11, 1).value = resolutions_text
        
    # 2. 清理第 15 行及之后的旧合并单元格
    existing_merged = [str(rng) for rng in ws.merged_cells.ranges if rng.min_row >= 15]
    for rng in existing_merged:
        ws.unmerge_cells(rng)
        
    font_default = Font(name="微软雅黑", size=10.0)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    start_row = 15
    for idx, item in enumerate(items, 1):
        r = start_row + idx - 1
        
        content = item.get("content") or item.get("待办项/未决事宜") or ""
        proposer = item.get("proposer") or item.get("提出人") or None
        owner = item.get("owner") or item.get("负责部门/人") or None
        deadline = item.get("deadline") or item.get("预计关闭时间") or None
        notes = item.get("notes") or item.get("备注") or None
        
        # 写入单元格值
        ws.cell(r, 1).value = idx
        ws.cell(r, 2).value = content
        ws.cell(r, 4).value = proposer
        ws.cell(r, 5).value = owner
        ws.cell(r, 6).value = deadline
        ws.cell(r, 7).value = notes
        
        # 合并 B 列与 C 列
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        
        # 统一字体、边框与对齐
        for c in range(1, 8):
            cell = ws.cell(r, c)
            cell.font = font_default
            
            l_side = "medium" if c == 1 else "thin"
            r_side = "medium" if c == 7 else "thin"
            if c == 3:
                l_side = None # 合并单元格内部
                
            cell.border = Border(
                left=Side(style=l_side, color="000000") if l_side else None,
                right=Side(style=r_side, color="000000") if r_side else None,
                top=Side(style="thin", color="000000"),
                bottom=Side(style="thin", color="000000")
            )
            
            if c in [1, 4, 5, 6]:
                cell.alignment = align_center
            elif c in [2, 7]:
                cell.alignment = align_left
                
        # 智能动态行高
        content_len = len(content) if content else 0
        notes_len = len(notes) if notes else 0
        max_len = max(content_len, notes_len)
        if max_len > 55:
            ws.row_dimensions[r].height = 42.0
        elif max_len > 28:
            ws.row_dimensions[r].height = 32.0
        else:
            ws.row_dimensions[r].height = 26.0

    save_target = output_path if output_path else excel_path
    wb.save(save_target)
    print(f"[AutoMeeting] 成功回填 {len(items)} 项待办事项至: {save_target}")
    return save_target

def main():
    parser = argparse.ArgumentParser(description="会议纪要 Excel 安全回填工具")
    parser.add_argument("-x", "--excel", required=True, help="目标 Excel 文件路径")
    parser.add_argument("-d", "--data", required=True, help="待办事项 JSON 文件路径")
    parser.add_argument("-r", "--resolutions", default=None, help="可选：更新决议内容的文本文件路径")
    parser.add_argument("-o", "--output", default=None, help="输出文件路径，若不填则原地更新原文件")
    
    args = parser.parse_args()
    try:
        with open(args.data, "r", encoding="utf-8") as f:
            items = json.load(f)
            
        res_text = None
        if args.resolutions and os.path.isfile(args.resolutions):
            with open(args.resolutions, "r", encoding="utf-8") as f:
                res_text = f.read()
                
        fill_minutes_excel(args.excel, items, res_text, args.output)
    except Exception as e:
        print(f"错误: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    main()
